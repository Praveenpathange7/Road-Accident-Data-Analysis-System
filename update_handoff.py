from docx import Document
from docx.oxml import OxmlElement
from pathlib import Path
path=Path('Road_Accident_Project_Master_Handoff.docx')
doc=Document(path)
def setcell(cell,text):
    p=cell.paragraphs[0]
    if p.runs:
        p.runs[0].text=text
        for r in list(p.runs[1:]): r._element.getparent().remove(r._element)
    else:
        p.add_run(text)
    for extra in list(cell.paragraphs[1:]): extra._element.getparent().remove(extra._element)
# Current status table
status=doc.tables[3]
setcell(status.rows[1].cells[1], 'Flask dashboard with filters, search, metrics, CSV/PDF reports and source links; MySQL schema and generated 2015-2024 national and State/UT aggregates; README, Render Blueprint and environment-variable setup.')
setcell(status.rows[2].cells[1], 'Provision an external MySQL host, import road_accident_database.sql, configure Render environment variables/network access, and deploy the web service. No public deployment URL exists until those account steps are completed. CSS comparisons are used instead of Matplotlib.')
setcell(status.rows[3].cells[1], 'The application is configured to require MySQL in production and to use CSV only for local development. The SQL builder was verified against the bundled 2022 State/UT snapshot; a live MySQL-host connection has not been tested.')
setcell(status.rows[4].cells[1], 'Push the project to GitHub without .env; provision MySQL 8 and import road_accident_database.sql; connect the repository in Render and enter database credentials in the Render environment settings; deploy and confirm /health is HTTP 200.')
# Team deliverables update for website owner
team=doc.tables[2]
setcell(team.rows[2].cells[2], 'Flask dashboard, production MySQL integration, local CSV mode, filters, CSV/PDF reports, resource links and Render deployment instructions.')
# Quick answers
faq=doc.tables[4]
setcell(faq.rows[2].cells[1], 'The SQL builder loads national annual totals for 2015-2024, State/UT accident counts for 2015-2024, and available State/UT casualty figures for 2017-2023. The 2022 snapshot is included in those series.')
setcell(faq.rows[4].cells[1], 'Yes. Production always queries MySQL and returns an explicit service error if the database is unavailable; CSV is available only as an optional local development source.')
# Starter context for another AI
p=doc.paragraphs[40]
old=p.text
status_start=old.find('Current project status:')
status_end=old.find('Ask focused questions', status_start)
if status_start < 0 or status_end < 0: raise SystemExit('Expected handoff status section was not found')
new=(old[:status_start]
    + 'Current project status: Flask app and database integration are implemented. Production requires MySQL; CSV is local-only. The SQL builder creates the MySQL schema and seeds the bundled 2015-2024 national/State/UT series. Hosting-provider provisioning, Render credentials/network access, and public deployment remain to be completed by the account owner. '
    + old[status_end:])
new=new.replace('The stack is MySQL, Flask, Pandas/NumPy and Matplotlib.', 'The stack is Flask, Pandas, MySQL Connector, MySQL 8, ReportLab and Gunicorn.')
new=new.replace('The current seed is official MoRTH/PIB state/UT aggregate data for 2022 only;', 'The bundled data includes national annual totals for 2015-2024 and State/UT accident totals for 2015-2024 with available casualty data for 2017-2023;')
if p.runs:
    p.runs[0].text=new
    for r in list(p.runs[1:]): r._element.getparent().remove(r._element)
else: p.add_run(new)
doc.save(path)
print(path.resolve())
